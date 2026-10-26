# -*- coding: utf-8 -*-
"""指示32 検証: 期待値突合 + pristineとのセル差分(数式テキスト基準)"""
import openpyxl, json, sys

sys.stdout.reconfigure(encoding='utf-8')
base = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot13"
EDIT = base + r"\比較グラフ_x1.3_original.xlsx"
PRIS = base + r"\比較グラフ_x1.3_pristine_20261026.xlsx"
exp = json.load(open(r"C:\Users\ss80751\kohden-code-workspace\20261019_鹿島引き込み_生産負荷シミュレーション\expected_values_20261026.json", encoding='utf-8'))

MKEY = exp['months']
L1_ROWS = [5, 6, 7, 8, 10, 11, 12]
L3 = {v: ([17, 18, 19, 20, 22, 23, 24] if v == 8 else [19, 20, 21, 22, 24, 25, 26]) for v in range(3, 9)}
L2ROWS = [31, 32, 33, 34, 36, 37, 38]  # 複製後のターゲット月行 (+2シフト)
errors = []

def chk(label, got, want, tol=1e-6):
    if got is None or abs(got - want) > tol * max(1.0, abs(want)):
        errors.append(f"!! {label}: got={got} want={want}")

# ---- 期待値突合 (data_only) ----
wbv = openpyxl.load_workbook(EDIT, data_only=True)
for v in (6, 7):
    ws = wbv[f'ver{v} (2)']
    chk(f'ver{v} (2) J21', ws.cell(21, 10).value, 19)
    chk(f'ver{v} (2) K21', ws.cell(21, 11).value, 14.75)
    chk(f'ver{v} (2) L21', ws.cell(21, 12).value, 280.25)
    chk(f'ver{v} (2) M21', ws.cell(21, 13).value, exp['expected'][f'ver{v}']['row21']['M'])
    chk(f'ver{v} (2) O7', ws.cell(7, 15).value, exp['expected'][f'ver{v}']['O']['2026-11']['O'])
for v in range(3, 8):
    ws = wbv[f'ver{v} (2)']
    e = exp['expected'][f'ver{v}']
    chk(f'ver{v} (2) I27', ws.cell(27, 9).value, e['I27'])
    for i, mk in enumerate(MKEY):
        r = L2ROWS[i]
        for c, item in ((2, 'TEC-5600'), (3, 'TEC-1000'), (4, '送信機_L2'), (5, 'L2停止時間')):
            chk(f'ver{v} (2) L2 {mk} {item}', ws.cell(r, c).value, exp['boss_sums']['3'][item][mk])
        chk(f'ver{v} (2) L2 L{r}', ws.cell(r, 12).value, 280.25)
    for i, mk in enumerate(MKEY):
        r1 = L1_ROWS[i]
        chk(f'ver{v} (2) O {mk}', ws.cell(r1, 15).value, e['O'][mk]['O'])
ws8 = wbv['ver8 (2)']
chk('ver8 (2) I13', ws8.cell(13, 9).value, exp['expected']['ver8']['I13'])
chk('ver8 (2) I25', ws8.cell(25, 9).value, exp['expected']['ver8']['I27'])
i37 = sum(exp['boss_sums']['8'][it][mk] for mk in MKEY for it in ('TEC-5600', 'TEC-1000', '送信機_L2', 'L2停止時間'))
chk('ver8 (2) I37', ws8.cell(37, 9).value, i37)
wsc = wbv['ver別簡易比較']
for j, v in enumerate(range(3, 9)):
    col = 3 + j
    e = exp['expected'][f'ver{v}']
    for i, mk in enumerate(MKEY):
        r = rowmap = [21, 22, 23, 24, 25, 26, 27][i]
        chk(f'簡易比較 {chr(64+col)}{r}', wsc.cell(r, col).value, e['O'][mk]['O'])
    chk(f'簡易比較 総計 col{col}', wsc.cell(28, col).value, e['total'])
b_avg = sum(wsc.cell(r, 2).value for r in range(21, 28)) / 7
chk('簡易比較 B28', wsc.cell(28, 2).value, b_avg)

# 構造確認
wbf = openpyxl.load_workbook(EDIT)
for v in range(3, 9):
    ws = wbf[f'ver{v} (2)']
    if len(ws.merged_cells.ranges):
        errors.append(f"!! ver{v} (2) 結合残存: {list(ws.merged_cells.ranges)}")
    ncf = len(list(ws.conditional_formatting))
    if ncf < 1:
        errors.append(f"!! ver{v} (2) CF無し")
total_pv = sum(len(ws._pivots) for ws in wbf.worksheets)
if total_pv != 26:
    errors.append(f"!! ピボット数 {total_pv} ≠ 26")

# ---- セル差分 (formulasモード) ----
wbe = openpyxl.load_workbook(EDIT)
wbp = openpyxl.load_workbook(PRIS)
diffs = []
for name in wbp.sheetnames:
    wse, wsp = wbe[name], wbp[name]
    maxr = max(wse.max_row, wsp.max_row)
    maxc = max(wse.max_column, wsp.max_column)
    for r in range(1, maxr + 1):
        for c in range(1, maxc + 1):
            ve, vp = wse.cell(r, c).value, wsp.cell(r, c).value
            if ve != vp:
                if ve is None and isinstance(vp, str) and vp.startswith('='):
                    ve_s = None
                diffs.append((name, wse.cell(r, c).coordinate, vp, ve))
print(f"差分セル数: {len(diffs)}")
from collections import Counter
cnt = Counter(d[0] for d in diffs)
print("シート別:", dict(cnt))
for d in diffs:
    is_block = d[0].endswith('(2)') and d[1].lstrip('ABCDEFGHIJKLMN').isdigit() and 29 <= int(d[1].lstrip('ABCDEFGHIJKLMN')) <= 39
    if not is_block:
        v_old = (str(d[2])[:40] if d[2] is not None else None)
        v_new = (str(d[3])[:40] if d[3] is not None else None)
        print(f"  {d[0]}!{d[1]}: {v_old} -> {v_new}")

# エラー値スキャン (全シートのキャッシュ値に #で始まるエラーがないか)
wbv2 = openpyxl.load_workbook(EDIT, data_only=True)
err_cells = []
for ws in wbv2.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('#') and c.value != '#':
                err_cells.append(f"{ws.title}!{c.coordinate}={c.value}")
print(f"エラー値セル: {len(err_cells)}件", err_cells[:10] if err_cells else "")

print()
if errors:
    print("検証エラー:")
    for e in errors:
        print(" ", e)
else:
    print("全期待値チェック PASS")
