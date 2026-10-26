# -*- coding: utf-8 -*-
"""ver別簡易比較の下段表(月別負荷率比較)からversion 2列のデータを削除
対象: B21:B32(月別静的値7つ+平均/最大/発生月/最小/発生月の数式5つ)。上段URS表とヘッダーは維持
方法: Excel COMでClearContents(書式維持)→再計算→保存→突合
"""
import openpyxl, shutil, sys
sys.stdout.reconfigure(encoding='utf-8')

BASE = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot13\比較グラフ_x1.3_平均ラベル_20261026.xlsx"
TMP = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot13\比較グラフ_x1.3_ver2削除_20261026.xlsx"
SHEET = 'ver別簡易比較'

# ---- Phase1: COMでB21:B32削除 ----
shutil.copyfile(BASE, TMP)
import win32com.client
app = win32com.client.DispatchEx("Excel.Application")
app.DisplayAlerts = False
app.Visible = False
try:
    wbx = app.Workbooks.Open(TMP)
    ws = wbx.Worksheets(SHEET)
    ws.Range("B21:B32").ClearContents()
    app.CalculateFullRebuild()
    wbx.Save()
    wbx.Close(SaveChanges=False)
    print("Phase1: B21:B32 ClearContents + 再計算 + 保存 完了")
finally:
    app.Quit()

# ---- Phase2: 突合 ----
def load(p, only):
    return openpyxl.load_workbook(p, data_only=only)
base_v, new_v = load(BASE, True), load(TMP, True)
base_f, new_f = load(BASE, False), load(TMP, False)
assert base_v.sheetnames == new_v.sheetnames, "シート構成の変化"
diffs = []
for sn in base_v.sheetnames:
    s1, s2 = base_v[sn], new_v[sn]
    mr, mc = max(s1.max_row, s2.max_row), max(s1.max_column, s2.max_column)
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            a, b = s1.cell(r, c).value, s2.cell(r, c).value
            if a != b:
                diffs.append((sn, s1.cell(r, c).coordinate, a, b))
vd = {(sn, co) for sn, co, _, _ in diffs}
expected = {(SHEET, f'B{r}') for r in range(21, 33)}
print(f"Phase2: 値差分 {len(diffs)}件(期待=12件 B21:B32)")
for d in diffs[:15]:
    print("   ", d[0], d[1], repr(d[2])[:40], '→', repr(d[3]))
fdiffs = []
for sn in base_f.sheetnames:
    s1, s2 = base_f[sn], new_f[sn]
    mr, mc = max(s1.max_row, s2.max_row), max(s1.max_column, s2.max_column)
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            a, b = s1.cell(r, c).value, s2.cell(r, c).value
            if a != b:
                fdiffs.append((sn, s1.cell(r, c).coordinate, a, b))
print(f"Phase2: 数式差分 {len(fdiffs)}件(期待=12件 B21:B32・静的値は両パスに出る)")
for d in fdiffs[:15]:
    print("   ", d[0], d[1], repr(d[2])[:40], '→', repr(d[3]))
npiv = sum(len(w._pivots) for w in new_f.worksheets)
errs = sum(1 for w in new_v.worksheets for row in w.iter_rows() for cell in row
           if isinstance(cell.value, str) and cell.value.startswith('#'))
wsv = new_v[SHEET]
wsf = new_f[SHEET]
keep_ok = (wsv['B20'].value == 'version 2' and wsv['B2'].value == 'version 2'
           and wsv['B3'].value == 'URS-4501-02' and wsv['B15'].value == 'URS-4527-01'
           and wsf['C21'].value == "='ver3 (2)'!O5" and wsv['A28'].value == '平均')
print(f"Phase2: ピボット {npiv}個 / エラー値 {errs}件 / 上段表・ヘッダー・他列の維持={keep_ok}")
exp_f = {(SHEET, f'B{r}') for r in range(21, 33)}
ok = (vd == expected and {(sn, co) for sn, co, _, _ in fdiffs} == exp_f
      and npiv == 26 and errs == 0 and keep_ok)
print("Phase2: 判定 =", "PASS" if ok else "NG")
sys.exit(0 if ok else 1)
