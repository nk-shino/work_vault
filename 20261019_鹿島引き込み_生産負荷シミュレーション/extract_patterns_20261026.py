# -*- coding: utf-8 -*-
"""指示32 パターン抽出: x1.4バックアップとx1.3比較ファイルの現状ダンプ(読み取り専用)"""
import openpyxl, sys
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

X14 = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot\backup_L2_20261020\比較グラフ_L2適用済み.xlsx"
X13 = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot13\比較グラフ_x1.3_original.xlsx"

def fmt_cell(cell):
    v = cell.value
    if v is None:
        return None
    s = repr(v)
    extra = []
    if cell.number_format != "General":
        extra.append(f"fmt={cell.number_format!r}")
    f = cell.font
    if f.bold:
        extra.append("B")
    if f.underline and f.underline != "none":
        extra.append(f"U={f.underline}")
    fill = cell.fill
    if fill.patternType:
        fg = fill.fgColor.rgb if fill.fgColor is not None else None
        bg = fill.bgColor.rgb if fill.bgColor is not None else None
        extra.append(f"fill={fill.patternType}/{fg}/{bg}")
    if extra:
        s += " <" + ",".join(extra) + ">"
    return s

def dump(wbv, wbf, name, r1, r2, c1, c2, title):
    ws, wf = wbv[name], wbf[name]
    print(f"\n--- {title} [{name}] {get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2} ---")
    for r in range(r1, r2 + 1):
        items = []
        for c in range(c1, c2 + 1):
            s = fmt_cell(ws.cell(r, c))
            fv = wf.cell(r, c).value
            if isinstance(fv, str) and fv.startswith("="):
                s = (s if s is not None else "None") + f" FORMULA={fv}"
            if s is not None:
                items.append(f"{get_column_letter(c)}{r}:{s}")
        if items:
            print("  " + " | ".join(items))

def cf_summary(wb, name):
    ws = wb[name]
    out = []
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            fill_desc = ""
            if rule.dxf is not None and rule.dxf.fill is not None:
                fg = rule.dxf.fill.fgColor.rgb if rule.dxf.fill.fgColor is not None else None
                bg = rule.dxf.fill.bgColor.rgb if rule.dxf.fill.bgColor is not None else None
                fill_desc = f"fill={fg}/{bg}"
            out.append(f"  {rng.sqref} type={rule.type} op={rule.operator} formula={rule.formula} {fill_desc} pri={rule.priority}")
    print(f"--- CF [{name}] ---")
    print("\n".join(out) if out else "  (none)")

print("=" * 20, "X14 (x1.4 backup)", "=" * 20)
wbf14 = openpyxl.load_workbook(X14, data_only=False)
wbv14 = openpyxl.load_workbook(X14, data_only=True)
print("sheets:", wbf14.sheetnames)

dump(wbv14, wbf14, 'ver別簡易比較', 18, 34, 1, 14, "簡易比較 現状")
cf_summary(wbv14, 'ver別簡易比較')

dump(wbv14, wbf14, 'ver3', 28, 41, 1, 19, "ver3 L2ブロック")
dump(wbv14, wbf14, 'ver3', 1, 28, 1, 19, "ver3 L1+L3")
cf_summary(wbv14, 'ver3')

dump(wbv14, wbf14, 'ver6', 15, 28, 1, 15, "ver6 L3(21行目の復元元)")
dump(wbv14, wbf14, 'ver8', 1, 28, 1, 19, "ver8 L1+L3")
dump(wbv14, wbf14, 'ver8', 28, 41, 1, 19, "ver8 L2")

print("\n" + "=" * 20, "X13 (x1.3 original)", "=" * 20)
wbf13 = openpyxl.load_workbook(X13, data_only=False)
wbv13 = openpyxl.load_workbook(X13, data_only=True)
print("sheets:", wbf13.sheetnames)

piv = {n: len(wbf13[n]._pivots) for n in wbf13.sheetnames if getattr(wbf13[n], '_pivots', None)}
print("pivot parts per sheet:", piv, "TOTAL:", sum(piv.values()))

targets = ['ver3 (2)', 'ver4 (2)', 'ver5 (2)', 'ver6 (2)', 'ver7 (2)', 'ver8 (2)', 'ver別簡易比較']
for t in targets:
    r2 = 34 if t == 'ver別簡易比較' else 30
    dump(wbv13, wbf13, t, 1, r2, 1, 19, t)
    cf_summary(wbv13, t)

dump(wbv13, wbf13, 'ver8 (2)', 26, 40, 1, 19, "ver8 (2) L2既存")
