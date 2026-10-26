# -*- coding: utf-8 -*-
"""指示32 最終監査: 納品ファイル(share)を第三者視点で再検証(読み取り専用)"""
import openpyxl, sys
from collections import Counter

sys.stdout.reconfigure(encoding='utf-8')
FINAL = r"\\nksv569.ad.nkc.co.jp\Share\mnt\DGS-SV\★個人フォルダ\篠原\サイクルタイムx1.3\_品目合計数別負荷率比較グラフ.xlsx"

L1_ROWS = [5, 6, 7, 8, 10, 11, 12]
MONTH_LABELS = ['9月', '10月', '11月', '12月', '1月', '2月', '3月']

wbv = openpyxl.load_workbook(FINAL, data_only=True)
wbf = openpyxl.load_workbook(FINAL)

print("=== [1] ブック概要 ===")
print("シート:", wbv.sheetnames)
total_pv = sum(len(ws._pivots) for ws in wbf.worksheets)
err = []
for ws in wbv.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('#') and len(c.value) > 1:
                err.append(f"{ws.title}!{c.coordinate}={c.value}")
print(f"ピボット定義: {total_pv} / エラー値セル: {len(err)} {err[:5]}")

print()
print("=== [2] (2)シート 構造・日本語表示 ===")
for v in range(3, 9):
    name = f'ver{v} (2)'
    wsv, wsf = wbv[name], wbf[name]
    l3h = 15 if v == 8 else 17                    # L3ヘッダ行
    l3r = [17, 18, 19, 20, 22, 23, 24] if v == 8 else [19, 20, 21, 22, 24, 25, 26]
    l2h = 27 if v == 8 else 29                    # L2ヘッダ行
    l2r = [l2h + 2, l2h + 3, l2h + 4, l2h + 5, l2h + 7, l2h + 8, l2h + 9]
    checks = []
    checks.append(("L1ヘッダA3+7項目(合計/)", wsv['A3'].value == '行ラベル' and all(isinstance(wsv.cell(3, c).value, str) and '合計 / ' in wsv.cell(3, c).value for c in (2, 3, 4, 5, 7, 8))))
    checks.append(("J3-K3-L3=稼働日/1日時間/保有工数", wsv['J3'].value == '稼働日' and wsv['K3'].value == '1日時間' and wsv['L3'].value == '保有工数'))
    checks.append(("L1総計13", wsv['A13'].value == '総計'))
    checks.append((f"L3ヘッダ{l3h}行(AED/BSM-1700L3/L3停止時間)", wsv.cell(l3h, 1).value == '行ラベル' and 'AED' in str(wsv.cell(l3h, 2).value) and 'BSM-1700L3' in str(wsv.cell(l3h, 3).value) and 'L3停止時間' in str(wsv.cell(l3h, 4).value)))
    checks.append(("L3総計", str(wsv.cell(l3r[-1] + 1, 1).value) == '総計'))
    checks.append((f"L2ブロック{l2h}行目ヘッダ(TEC-5600等)", wsv.cell(l2h, 1).value == '行ラベル' and 'TEC-5600' in str(wsv.cell(l2h, 2).value) and 'L2停止時間' in str(wsv.cell(l2h, 5).value)))
    checks.append(("L2月7行", all(isinstance(wsv.cell(r, 1).value, str) for r in l2r)))
    checks.append(("L2総計+L2ラベルM列", str(wsv.cell(l2h + 10, 1).value) == '総計' and str(wsv.cell(l2h, 13).value) == 'L2'))
    checks.append(("N列注記(新規追分)", any('新規追加分' in str(wsv.cell(r, 14).value or '') for r in (l2h,))))
    mrow = 3 if str(wsv.cell(3, 13).value) == 'L1' else 4
    checks.append((f"L1ラベルM{mrow}+O列ラベル", str(wsv.cell(mrow, 13).value) == 'L1' and 'L1+L3' in str(wsv.cell(mrow, 15).value)))
    l3lab = 15 if v == 8 else 18
    checks.append((f"L3ラベルM{l3lab}", 'L3' in str(wsv.cell(l3lab, 13).value)))
    checks.append(("結合0", len(wsf.merged_cells.ranges) == 0))
    if v in (6, 7):
        checks.append(("行21復元 J/K/L/M", wsv.cell(21, 10).value == 19 and abs(wsv.cell(21, 12).value - 280.25) < 1e-9 and abs(wsv.cell(21, 13).value - 184.36446502976187 / 280.25) < 1e-9))
    ng = [k for k, ok in checks if not ok]
    print(f"  {name}: {'PASS' if not ng else 'NG ' + str(ng)}  (L2先頭値={wsv.cell(l2h + 2, 2).value})")

print()
print("=== [3] 書式(数値表示形式) ===")
FMT_I = '#,##0.0;[Red]\\-#,##0.0'
for v in range(3, 9):
    name = f'ver{v} (2)'
    wsf = wbf[name]
    l3r = [17, 18, 19, 20, 22, 23, 24] if v == 8 else [19, 20, 21, 22, 24, 25, 26]
    rows_i = L1_ROWS + [13] + l3r + [l3r[-1] + 1]
    fI = {wsf.cell(r, 9).number_format for r in rows_i}
    fM = {wsf.cell(r, 13).number_format for r in rows_i}
    fO = {wsf.cell(r, 15).number_format for r in L1_ROWS + [13]}
    krows = [10, 11, 12] + ([24, 25, 26] if v != 8 else [])
    fK = {wsf.cell(r, 11).number_format for r in krows}
    okI = fI <= {FMT_I}
    okM = fM <= {'0%'}
    okO = fO <= {'0%'}
    okK = fK <= {'0.00"h"'}
    print(f"  {name}: I={'OK' if okI else fI} M={'OK' if okM else fM} O={'OK' if okO else fO} K={'OK' if okK else fK}")

print()
print("=== [4] 条件付き書式(強調表示) ===")
def cf_info(wsf):
    out = []
    for rng in wsf.conditional_formatting:
        for rule in rng.rules:
            fill = None
            try:
                fill = rule.dxf.fill.bgColor.rgb or rule.dxf.fill.fgColor.rgb
            except Exception:
                pass
            op = getattr(rule, 'operator', None)
            fml = rule.formula if hasattr(rule, 'formula') else None
            out.append((str(rng.sqref), op, fml, fill))
    return out

for v in range(3, 9):
    name = f'ver{v} (2)'
    wsf, wsv = wbf[name], wbv[name]
    infos = cf_info(wsf)
    # 強調されるセル(値>=2)をCF範囲内で実数
    hi = []
    l3r = [17, 18, 19, 20, 22, 23, 24] if v == 8 else [19, 20, 21, 22, 24, 25, 26]
    l2f = 27 if v == 8 else 29
    rowsM = L1_ROWS + l3r + [l2f + 2, l2f + 3, l2f + 4, l2f + 5, l2f + 7, l2f + 8, l2f + 9]
    for r in rowsM:
        val = wsv.cell(r, 13).value
        if isinstance(val, (int, float)) and val >= 2:
            hi.append(f"M{r}({wsv.cell(r,1).value})={val:.2f}")
    for r in L1_ROWS:
        val = wsv.cell(r, 15).value
        if isinstance(val, (int, float)) and val >= 2:
            hi.append(f"O{r}({wsv.cell(r,1).value})={val:.2f}")
    sq = infos[0][0] if infos else 'なし'
    print(f"  {name}: CF範囲={sq[:70]}... 条件={infos[0][1] if infos else '-'} 塗={infos[0][3] if infos else '-'}")
    print(f"    強調セル(200%以上): {len(hi)}件 {hi[:6]}")

print()
print("=== [5] ピボットテーブル(定義と値の存在) ===")
from openpyxl.utils import range_boundaries
n_pv = 0
for name in wbf.sheetnames:
    wsf, wsv = wbf[name], wbv[name]
    for pt in wsf._pivots:
        n_pv += 1
        try:
            ref = pt.location.ref
            dfields = [df.name for df in pt.dataFields]
            rfields = []
            try:
                for rf in pt.rowFields:
                    if rf.x == -2:
                        rfields.append('値')
                    else:
                        rfields.append(pt.cache.cacheFields[rf.x].name)
            except Exception:
                rfields = ['?']
            c1, r1, c2, r2 = range_boundaries(ref)
            nnum = sum(1 for rr in wsv.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2) for cc in rr if isinstance(cc.value, (int, float)))
            src = getattr(getattr(pt.cache, 'cacheSource', None), 'ref', '?') or getattr(getattr(pt.cache.cacheSource, 'worksheetSource', None), 'ref', '?')
            print(f"  {name}: {ref} 行={rfields} 値={len(dfields)}項目 数値セル{nnum}個 src={src}")
        except Exception as e:
            print(f"  {name}: 定義読取失敗 {e}")
print(f"  合計 {n_pv}定義")

print()
print("=== [6] 簡易比較シート ===")
wsv, wsf = wbv['ver別簡易比較'], wbf['ver別簡易比較']
title = None
for r in range(1, 21):
    for c in range(1, 6):
        val = wsv.cell(r, c).value
        if isinstance(val, str) and '合計負荷率比較' in val:
            title = (wsv.cell(r, c).coordinate, val)
print(f"  タイトル: {title}")
print(f"  20行ヘッダ: {[wsv.cell(20, c).value for c in range(1, 9)]}")
print(f"  19行ヘッダ: {[wsv.cell(19, c).value for c in range(1, 9)]}")
print(f"  A21:A28: {[wsv.cell(r, 1).value for r in range(21, 29)]}")
print(f"  B21:B27(ユーザー値): {[round(wsv.cell(r, 2).value, 4) for r in range(21, 28)]}")
print(f"  C21式例: {wsf['C21'].value} → 値={wsv['C21'].value:.4f}")
print(f"  H27式例: {wsf['H27'].value} → 値={wsv['H27'].value:.4f}")
print(f"  28行(総計): 値={[round(wsv.cell(28, c).value, 4) for c in range(2, 9)]} 太字={all(wsf.cell(28, c).font.bold for c in range(2, 9))} 書式={wsf['C28'].number_format}")
print(f"  29/30行: {[wsv.cell(29, c).value for c in range(1, 9)]} / {[wsv.cell(30, c).value for c in range(1, 9)]}")
fC28 = wsf['C28'].number_format
print(f"  B28書式={wsf['B28'].number_format} 太字={wsf['B28'].font.bold}")
infos = cf_info(wsf)
hi = []
for c in range(3, 9):
    for r in range(21, 29):
        val = wsv.cell(r, c).value
        if isinstance(val, (int, float)) and val >= 2:
            hi.append(f"{wsv.cell(20, c).value}/{wsv.cell(r, 1).value}={val:.2f}")
print(f"  CF: {[(str(i[0]), i[1], i[3]) for i in infos]}")
print(f"  強調セル(200%以上): {len(hi)}件 {hi}")

print()
print("=== [7] 日本語・文字化けスキャン ===")
bad = []
jp_keywords = {'行ラベル': 0, '総計': 0, '最大': 0, '└発生月': 0, '新規追加分': 0, '2026年': 0, '2027年': 0}
for ws in wbv.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                if '\ufffd' in c.value or any(ord(ch) < 32 and ch not in '\n\t\r' for ch in c.value):
                    bad.append(f"{ws.title}!{c.coordinate}={c.value[:20]}")
                if ws.title.endswith('(2)') or ws.title == 'ver別簡易比較':
                    for k in jp_keywords:
                        if k in c.value:
                            jp_keywords[k] += 1
print(f"  文字化け候補: {len(bad)}件 {bad[:5]}")
print(f"  主要日本語ラベル出現数: {jp_keywords}")
