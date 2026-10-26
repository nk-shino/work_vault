# -*- coding: utf-8 -*-
"""指示32 検算基礎: BOSS x1.3計算シートの月次再集計と(2)シート・簡易比較の突合(読み取り専用)"""
import openpyxl, sys, json
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

X14 = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot\backup_L2_20261020\比較グラフ_L2適用済み.xlsx"
X13 = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot13\比較グラフ_x1.3_original.xlsx"
BOSS_DIR = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot13"
BOSS = {v: rf"{BOSS_DIR}\BOSS計画_鹿島4(半製品が鹿島除く)+NKV_1 (version {v}) (1).xlsm" for v in range(3, 9)}

CAL = {'フリー': 6, 'L1試作': 7, 'NKV': 8, 'NKV-440': 9, 'NKV-550': 10, 'BSM-1700': 11,
       'L1停止時間': 12, 'TEC-5600': 14, 'TEC-1000': 15, '送信機_L2': 16, 'L2停止時間': 17,
       'AED': 19, 'BSM-1700L3': 20, 'L3停止時間': 21}
MONTHS = [(2026, 9), (2026, 10), (2026, 11), (2026, 12), (2027, 1), (2027, 2), (2027, 3)]
MKEY = ['2026-09', '2026-10', '2026-11', '2026-12', '2027-01', '2027-02', '2027-03']
MJ = {'2026-09': 19, '2026-10': 18, '2026-11': 19, '2026-12': 19, '2027-01': 18, '2027-02': 18, '2027-03': 18}
KC = 14.75
LCAP = sum(MJ[m] for m in MKEY) * KC  # 1902.75

# ---- 1. BOSS x1.3 計算シートの月次合計 ----
boss_sums = {}
bsml3_nonempty = 0
for v, p in BOSS.items():
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['計算']
    sums = {k: defaultdict(float) for k in CAL}
    cnt = {k: defaultdict(int) for k in CAL}
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if not hasattr(d, 'year'):
            continue
        key = (d.year, d.month)
        for name, c in CAL.items():
            val = ws.cell(r, c).value
            if isinstance(val, (int, float)):
                sums[name][key] += val
                cnt[name][key] += 1
    boss_sums[v] = {k: {mk: sums[k][(y, m)] for mk, (y, m) in zip(MKEY, MONTHS)} for k in CAL}
    bsml3_nonempty += sum(cnt['BSM-1700L3'].values())
    wb.close()
print(f"[1] BOSS 6ファイル月次集計完了。BSM-1700L3の非空セル数={bsml3_nonempty}")

l2items = ['TEC-5600', 'TEC-1000', '送信機_L2', 'L2停止時間']
ok = True
for it in l2items:
    for v in range(4, 9):
        for mk in MKEY:
            if abs(boss_sums[v][it][mk] - boss_sums[3][it][mk]) > 1e-9:
                ok = False
                print(f"  !! L2不一致 ver{v} {it} {mk}")
print(f"[2] L2項目のバージョン共通性: {'OK(全ver一致)' if ok else 'NG'}")

# ---- 2. x1.4 簡易比較の総計(参考情報) ----
wb14 = openpyxl.load_workbook(X14)
ws14 = wb14['ver2']
def rsum(ws, r, c1=2, c2=8):
    return sum(v for c in range(c1, c2 + 1) if isinstance((v := ws.cell(r, c).value), (int, float)))
I13_14, I27_14, I39_14 = rsum(ws14, 13), rsum(ws14, 27), rsum(ws14, 39)
wsK = wb14['ver別簡易比較']
B28 = wsK.cell(28, 2).value
o14 = [wsK.cell(r, 2).value for r in range(21, 28)]
print(f"[3] 参考 x1.4: I13={I13_14:.2f} I27={I27_14:.2f} I39={I39_14:.2f} / 加重={(I13_14+I27_14+I39_14)/LCAP:.6f} / 月別平均={sum(o14)/7:.6f} / B28={B28}")
wb14.close()

# ---- 3. (2)シート ピボット本体 vs BOSS計算 突合 ----
wb13 = openpyxl.load_workbook(X13, data_only=True)
L1_ROWS = [5, 6, 7, 8, 10, 11, 12]
L3_ROWS_STD = [19, 20, 21, 22, 24, 25, 26]
L3_ROWS_V8 = [17, 18, 19, 20, 22, 23, 24]

def close(a, b):
    if a is None:
        return abs(b) <= 1e-9
    return abs(a - b) <= 1e-6 * max(1.0, abs(b))

mismatch = 0
checked = 0
expected = {}
for ver in range(3, 9):
    name = f'ver{ver} (2)'
    ws = wb13[name]
    l3rows = L3_ROWS_V8 if ver == 8 else L3_ROWS_STD
    exp = {'O': {}, 'row21': None}
    hdr = {c: str(ws.cell(3, c).value or '').replace('合計 / ', '').replace('個数 / ', '') for c in range(2, 9)}
    for i, mk in enumerate(MKEY):
        r = L1_ROWS[i]
        for c in range(2, 9):
            item = hdr[c]
            if item in ('', 'None'):
                continue
            checked += 1
            if not close(ws.cell(r, c).value, boss_sums[ver][item][mk]):
                mismatch += 1
                print(f"  !! {name} L1 {ws.cell(r,1).value} {item}: pivot={ws.cell(r,c).value} calc={boss_sums[ver][item][mk]}")
    hr = l3rows[0] - 2
    hdr3 = {c: str(ws.cell(hr, c).value or '').replace('合計 / ', '').replace('個数 / ', '') for c in range(2, 5)}
    for i, mk in enumerate(MKEY):
        r = l3rows[i]
        for c in range(2, 5):
            item = hdr3[c]
            if item in ('', 'None'):
                continue
            checked += 1
            if not close(ws.cell(r, c).value, boss_sums[ver][item][mk]):
                mismatch += 1
                print(f"  !! {name} L3 {ws.cell(r,1).value} {item}: pivot={ws.cell(r,c).value} calc={boss_sums[ver][item][mk]}")
    prob = []
    for i, mk in enumerate(MKEY):
        r1, r3 = L1_ROWS[i], l3rows[i]
        I1, L1v, M1 = ws.cell(r1, 9).value, ws.cell(r1, 12).value, ws.cell(r1, 13).value
        I3, L3v, M3 = ws.cell(r3, 9).value, ws.cell(r3, 12).value, ws.cell(r3, 13).value
        skip3 = (ver in (6, 7) and r3 == 21)  # 復元対象行はJ/K/L/Mが欠落
        if not skip3:
            J1v, K1v = ws.cell(r1, 10).value, ws.cell(r1, 11).value
            J3v, K3v = ws.cell(r3, 10).value, ws.cell(r3, 11).value
            if not close(L1v, J1v * K1v): prob.append(f"L1 L{r1}={L1v}≠J*K={J1v*K1v}")
            if not close(L3v, J3v * K3v): prob.append(f"L3 L{r3}={L3v}≠J*K={J3v*K3v}")
            if I1 is not None and L1v and not close(M1, I1 / L1v): prob.append(f"L1 M{r1}")
            if I3 is not None and L3v and not close(M3, I3 / L3v): prob.append(f"L3 M{r3}")
        exp['O'][mk] = {'M1': M1, 'M3': M3, 'O': ws.cell(r1, 15).value}
    if ver in (6, 7):
        I21 = ws.cell(21, 9).value
        m21 = I21 / 280.25
        exp['O']['2026-11']['M3'] = m21
        exp['O']['2026-11']['O'] = exp['O']['2026-11']['M1'] + m21
        exp['row21'] = {'J': 19, 'K': 14.75, 'I': I21, 'L': 280.25, 'M': m21}
    i13 = sum(ws.cell(r, 9).value for r in L1_ROWS)          # 総計行にIセルが無いver8 (2)対策→月行の合計
    i27 = sum(ws.cell(r, 9).value for r in l3rows)           # L3総計行にIセルは無い→月行の合計
    if ws.cell(13, 9).value is not None and not close(ws.cell(13, 9).value, i13):
        prob.append(f"I13キャッシュ{ws.cell(13, 9).value}≠月行合計{i13}")
    ovals = [exp['O'][mk]['O'] for mk in MKEY]
    exp['I13'], exp['I27'] = i13, i27
    exp['total'] = sum(ovals) / 7          # 総計=月別平均(x1.4のB28が平均に近い値のため)
    exp['total_weighted'] = (i13 + i27) / LCAP  # 参考: 加重値
    expected[f'ver{ver}'] = exp
    if prob:
        print(f"  !! {name} 内部整合: {prob}")
print(f"[4] (2)シート ピボット本体 vs BOSS計算: 比較{checked}件、不一致{mismatch}件")

ws8 = wb13['ver8 (2)']
hdr2 = {c: str(ws8.cell(27, c).value or '').replace('合計 / ', '') for c in range(2, 6)}
mm = 0
for i, mk in enumerate(MKEY):
    r = [29, 30, 31, 32, 34, 35, 36][i]
    for c in range(2, 6):
        item = hdr2[c]
        if item in ('', 'None'):
            continue
        if not close(ws8.cell(r, c).value, boss_sums[8][item][mk]):
            mm += 1
            print(f"  !! ver8 (2) L2 {ws8.cell(r,1).value} {item}: block={ws8.cell(r,c).value} calc={boss_sums[8][item][mk]}")
print(f"[5] ver8 (2) L2ブロック vs BOSS計算: 不一致{mm}件")

wsc = wb13['ver別簡易比較']
b_ok = all(close(wsc.cell(21 + i, 2).value, expected['ver3']['O'][mk]['O']) for i, mk in enumerate(MKEY))
print(f"[6] 簡易比較B21:B27 == ver3 (2) O列: {'完全一致' if b_ok else '不一致'}")

out = {'months': MKEY, 'J': MJ, 'K': KC, 'LCAP': LCAP,
       'boss_sums': {str(v): boss_sums[v] for v in boss_sums},
       'expected': {k: {'O': e['O'], 'I13': e['I13'], 'I27': e['I27'], 'total': e['total'],
                        'total_weighted': e['total_weighted'], 'row21': e['row21']} for k, e in expected.items()}}
with open(r"C:\Users\ss80751\kohden-code-workspace\20261019_鹿島引き込み_生産負荷シミュレーション\expected_values_20261026.json", 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=1, default=str)
print("[7] 期待値JSON保存完了")
for ver in range(3, 9):
    e = expected[f'ver{ver}']
    print(f"ver{ver}: " + " ".join(f"{mk[-2:]}月O={e['O'][mk]['O']:.4f}" for mk in MKEY) + f" 総計={e['total']:.4f}")
