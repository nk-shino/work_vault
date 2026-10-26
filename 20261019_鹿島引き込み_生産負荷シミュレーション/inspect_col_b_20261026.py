# -*- coding: utf-8 -*-
"""ver別簡易比較シートの削除前点検(読み取り専用): 使用範囲・結合・CF・図形・他シートからの参照"""
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

P = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot13\比較グラフ_x1.3_平均ラベル_20261026.xlsx"
wb = openpyxl.load_workbook(P, data_only=False)
ws = wb['ver別簡易比較']
print(f"使用範囲: {ws.max_row}行 × {ws.max_column}列")
print(f"結合セル: {[str(m) for m in ws.merged_cells.ranges]}")
print(f"図表: charts={len(getattr(ws, '_charts', []))} images={len(getattr(ws, '_images', []))}")
print("条件付き書式:")
for rng in ws.conditional_formatting:
    for rule in rng.rules:
        print(f"  {rng.sqref} {rule.type} {getattr(rule, 'operator', '')} {getattr(rule, 'formula', '')}")
print("\nセル内容 A1:L32(値 or 数式):")
for r in range(1, 33):
    row = []
    for c in range(1, 13):
        v = ws.cell(r, c).value
        if v is None:
            row.append('·')
        else:
            s = str(v)
            row.append(s[:24])
    if any(x != '·' for x in row):
        print(f"  {r:2d} | " + " | ".join(row))
print("\n他シートから'ver別簡易比較'への参照:")
hits = 0
for sn in wb.sheetnames:
    if sn == 'ver別簡易比較':
        continue
    s = wb[sn]
    for row in s.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and 'ver別簡易比較' in cell.value:
                print(f"  {sn}!{cell.coordinate}: {cell.value[:60]}")
                hits += 1
print(f"  参照 {hits}件")
