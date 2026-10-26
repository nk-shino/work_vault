# -*- coding: utf-8 -*-
"""BOSS version2 vs version3: 計算シートのデータ差分とL1/L3月次合計の比較(読み取り専用)"""
import openpyxl, sys
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')
BASE = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot\backup_L2_20261020"
V2 = BASE + r"\BOSS_v2_L2適用済み.xlsm"
V3 = BASE + r"\BOSS_v3_L2適用済み.xlsm"

w2 = openpyxl.load_workbook(V2, data_only=True)
w3 = openpyxl.load_workbook(V3, data_only=True)
print("シート構成 v2:", w2.sheetnames)
print("シート構成 v3:", w3.sheetnames)

s2, s3 = w2['計算'], w3['計算']
print(f"計算シート: v2={s2.max_row}行×{s2.max_column}列 / v3={s3.max_row}行×{s3.max_column}列")

diff = 0
maxr = max(s2.max_row, s3.max_row)
maxc = max(s2.max_column, s3.max_column)
for r in range(1, maxr + 1):
    for c in range(1, maxc + 1):
        a, b = s2.cell(r, c).value, s3.cell(r, c).value
        if a != b:
            if isinstance(a, float) and isinstance(b, float) and abs(a - b) < 1e-9:
                continue
            diff += 1
            if diff <= 10:
                print(f"  差分 R{r}C{c}: v2={a} v3={b}")
print(f"計算シートのセル差分: {diff}件")

# L1/L3項目の月次合計比較
CAL = {'フリー': 6, 'L1試作': 7, 'NKV': 8, 'NKV-440': 9, 'NKV-550': 10, 'BSM-1700': 11,
       'L1停止時間': 12, 'AED': 19, 'BSM-1700L3': 20, 'L3停止時間': 21}
MONTHS = [(2026, 9), (2026, 10), (2026, 11), (2026, 12), (2027, 1), (2027, 2), (2027, 3)]

def sums(ws):
    out = {k: defaultdict(float) for k in CAL}
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if not hasattr(d, 'year'):
            continue
        for k, c in CAL.items():
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                out[k][(d.year, d.month)] += v
    return out

m2, m3 = sums(s2), sums(s3)
mism = 0
for k in CAL:
    for mk in MONTHS:
        if abs(m2[k][mk] - m3[k][mk]) > 1e-9:
            mism += 1
            print(f"  月次不一致 {k} {mk}: v2={m2[k][mk]:.4f} v3={m3[k][mk]:.4f}")
print(f"L1/L3月次合計の不一致: {mism}件(比較{len(CAL)*7}組)")
