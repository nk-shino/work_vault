# -*- coding: utf-8 -*-
"""指示32 実装: (2)シートの修正と簡易比較の埋め (openpyxl・1回のみ実行)"""
import openpyxl, re, json, sys
from copy import copy
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill, Font

sys.stdout.reconfigure(encoding='utf-8')
X13 = r"C:\Users\ss80751\AppData\Local\Temp\kohden-code\boss_pivot13\比較グラフ_x1.3_original.xlsx"

L1_ROWS = [5, 6, 7, 8, 10, 11, 12]
L3 = {v: ([17, 18, 19, 20, 22, 23, 24] if v == 8 else [19, 20, 21, 22, 24, 25, 26]) for v in range(3, 9)}
FMT_I = '#,##0.0;[Red]\\-#,##0.0'
FMT_PCT = '0%'
FMT_PCT1 = '0.0%'
FMT_H = '0.00"h"'

wb = openpyxl.load_workbook(X13)
log = []

def set_bold(ws, coord, on=True):
    c = ws[coord]
    f = copy(c.font)
    f.bold = on
    c.font = f

# ---- ⓪ 結合セル全解除 (x1.4には結合なし・I21:M21がJ-M値を破壊済み。トップレフト値は保持) ----
for v in range(3, 9):
    ws = wb[f'ver{v} (2)']
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
        log.append(f"⓪ ver{v} (2): 結合解除 {rng}")

# ---- ① ver6/7 (2) 行21復元 (J21/K21/L21/M21) ----
for v in (6, 7):
    ws = wb[f'ver{v} (2)']
    vals = {10: 19, 11: 14.75, 12: '=J21*K21', 13: '=I21/L21'}
    for col, val in vals.items():
        src, dst = ws.cell(20, col), ws.cell(21, col)
        dst.value = val
        dst._style = copy(src._style)
    log.append(f"① ver{v} (2): J21=19 K21=14.75 L21==J21*K21 M21==I21/L21 復元(スタイルは20行目から)")

# ---- ② 総計行のIセル補完 (ver8 (2) I13 / 各シートL3総計I / ver8 (2) L2総計I) ----
ws8 = wb['ver8 (2)']
if ws8.cell(13, 9).value is None:
    ws8.cell(13, 9).value = '=SUM(B13:H13)'
    ws8.cell(13, 9)._style = copy(ws8.cell(12, 9)._style)
    log.append("② ver8 (2): I13==SUM(B13:H13) 補完")
for v in range(3, 8):
    ws = wb[f'ver{v} (2)']
    r3t = L3[v][-1] + 1  # 27
    if ws.cell(r3t, 9).value is None:
        ws.cell(r3t, 9).value = f'=SUM(B{r3t}:H{r3t})'
        ws.cell(r3t, 9)._style = copy(ws.cell(13, 9)._style)
        log.append(f"② ver{v} (2): I{r3t}==SUM(B{r3t}:H{r3t}) 補完")
if ws8.cell(25, 9).value is None:
    ws8.cell(25, 9).value = '=SUM(B25:H25)'
    ws8.cell(25, 9)._style = copy(ws8.cell(13, 9)._style)
    log.append("② ver8 (2): I25==SUM(B25:H25) 補完")
if ws8.cell(37, 9).value is None:
    ws8.cell(37, 9).value = '=SUM(B37:H37)'
    ws8.cell(37, 9)._style = copy(ws8.cell(13, 9)._style)
    log.append("② ver8 (2): I37==SUM(B37:H37) 補完(L2総計)")

# ---- ③ L2ブロック複製: ver8 (2) 27-37行 → ver3-7 (2) 29-39行 (+2シフト) ----
def shift_formula(f, delta=2):
    return re.sub(r'(\$?[A-Z]{1,3})(\$?)(\d{1,3})',
                  lambda m: f"{m.group(1)}{m.group(2)}{int(m.group(3)) + delta}", f)

for v in range(3, 8):
    ws = wb[f'ver{v} (2)']
    n_f = 0
    for sr in range(27, 38):
        tr = sr + 2
        for c in range(1, 15):  # A..N
            s, d = ws8.cell(sr, c), ws.cell(tr, c)
            val = s.value
            if isinstance(val, str) and val.startswith('='):
                val = shift_formula(val)
                n_f += 1
            d.value = val
            d._style = copy(s._style)
        h = ws8.row_dimensions[sr].height
        if h is not None:
            ws.row_dimensions[tr].height = h
    log.append(f"③ ver{v} (2): L2ブロック 29-39行へ複製(数式{n_f}件シフト)")

# ---- ④ 条件付き書式: M(L1+L3+L2)とOに @>=2 ----
fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
for v in range(3, 9):
    ws = wb[f'ver{v} (2)']
    if v == 8:
        mr = ['M5:M8', 'M10:M12', 'M17:M20', 'M22:M24', 'M29:M32', 'M34:M36']
    else:
        mr = ['M5:M8', 'M10:M12', 'M19:M22', 'M24:M26', 'M31:M34', 'M36:M38']
    sqref = ' '.join(mr + ['O5:O8', 'O10:O12'])
    ws.conditional_formatting.add(sqref, CellIsRule(operator='greaterThanOrEqual', formula=['2'], fill=fill))
    log.append(f"④ ver{v} (2): CF追加 {sqref} @>=2")

# ---- ⑤ 書式統一 ----
for v in range(3, 8):  # L3 K列 K24:K26
    ws = wb[f'ver{v} (2)']
    for r in (24, 25, 26):
        ws.cell(r, 11).number_format = FMT_H
log.append("⑤ ver3-7 (2): L3 K24:K26 に 0.00\"h\"")
for v in (3, 4, 5):  # L1 K列 K10:K12
    ws = wb[f'ver{v} (2)']
    for r in (10, 11, 12):
        ws.cell(r, 11).number_format = FMT_H
log.append("⑤ ver3-5 (2): L1 K10:K12 に 0.00\"h\"")
ws7 = wb['ver7 (2)']
ws7.cell(26, 13).number_format = FMT_PCT
for coord in ('J25', 'K25', 'L25', 'M25'):
    set_bold(ws7, coord, False)
log.append("⑤ ver7 (2): M26 に 0% / J25:K25:L25:M25 の太字解除")
for v in (3, 4, 5):
    ws = wb[f'ver{v} (2)']
    for coord in ('J3', 'K3', 'L3', 'M4', 'O4', 'M18'):
        set_bold(ws, coord, True)
log.append("⑤ ver3-5 (2): J3:L3/M4/O4/M18 太字")
for coord in ('M3', 'O3', 'M18'):
    set_bold(ws7, coord, True)
log.append("⑤ ver7 (2): M3/O3/M18 太字")
ws6 = wb['ver6 (2)']
set_bold(ws6, 'M18', True)
log.append("⑤ ver6 (2): M18 太字")

# I/M/O 標準書式(冪等)
for v in range(3, 9):
    ws = wb[f'ver{v} (2)']
    rows = L1_ROWS + [13] + L3[v] + [L3[v][-1] + 1]
    for r in rows:
        ws.cell(r, 9).number_format = FMT_I
        ws.cell(r, 13).number_format = FMT_PCT
    for r in L1_ROWS + [13]:
        ws.cell(r, 15).number_format = FMT_PCT
log.append("⑤ ver3-8 (2): I/M/O 標準書式 冪等適用")

# ---- ⑥ 簡易比較 C21:H27 + 総計行28 ----
wsc = wb['ver別簡易比較']
colmap = {3: 'ver3 (2)', 4: 'ver4 (2)', 5: 'ver5 (2)', 6: 'ver6 (2)', 7: 'ver7 (2)', 8: 'ver8 (2)'}
rowmap = {21: 5, 22: 6, 23: 7, 24: 8, 25: 10, 26: 11, 27: 12}
for c, sname in colmap.items():
    for r, orow in rowmap.items():
        d = wsc.cell(r, c)
        d.value = f"='{sname}'!O{orow}"
        d._style = copy(wsc.cell(21, 2)._style)  # B21のスタイルを基準に
log.append("⑥ 簡易比較 C21:H27 = 'verN (2)'!O行 (スタイルはB21から)")
wsc.cell(28, 2).value = '=AVERAGE(B21:B27)'
for c, sname in colmap.items():
    wsc.cell(28, c).value = f"=AVERAGE('{sname}'!O5:O8,'{sname}'!O10:O12)"
    wsc.cell(28, c)._style = copy(wsc.cell(21, 2)._style)
for c in range(2, 9):
    cell = wsc.cell(28, c)
    cell.number_format = FMT_PCT1
    f = copy(cell.font)
    f.bold = True
    cell.font = f
log.append("⑥ 簡易比較 28行 B:H = 月別平均 (太字 0.0%)")

wb.save(X13)
print("保存完了:", X13)
for line in log:
    print(" -", line)

# ピボット部品数の確認
total_pv = sum(len(ws._pivots) for ws in wb.worksheets)
print(f"ピボット部品総数: {total_pv} (期待52)")
