# BOSS計画_鹿島4(version 3) の chkTime/Main 時間負荷計算をPythonで再現する
# 読み取り専用。元ファイルは変更しない。既存計算の検証と再計算の基礎。
import openpyxl, sys, re
from collections import defaultdict
from datetime import datetime, date

P = r"C:\Users\ss80751\Downloads\BOSS計画_鹿島4(半製品が鹿島除く)+NKV_1 (version 3).xlsm"
sys.stdout.reconfigure(encoding='utf-8')

wbf = openpyxl.load_workbook(P, data_only=False)
wbv = openpyxl.load_workbook(P, data_only=True)

# ---------- VBA Main / chkTime の再現 ----------
KARI = 56          # 量産 CT代替(秒/台)
KARI_SISAKU = 75   # 試作 CT代替(秒/台)
SETUP = 600        # 段取(秒/指図)
EFF = 0.7          # ÷0.7

# D_CycleTime: サイクルタイムシート A列→D列(初回出現優先)
D_CycleTime = {}
ws = wbv['サイクルタイム']
for r in range(2, ws.max_row+1):
    a = ws.cell(r,1).value
    if a is None or str(a).strip()=="":
        continue
    a = str(a)
    if a in D_CycleTime:
        continue  # 上の行が優先(後続は警告のみ)
    v = ws.cell(r,4).value
    D_CycleTime[a] = v  # 数値以外は後段でVarTypeチェック相当の扱いにする

# D_Sisaku: 試作シート A→(C固定生産時間, D組取り数)
D_Sisaku = {}
ws = wbv['試作']
for r in range(2, ws.max_row+1):
    a = ws.cell(r,1).value
    if a is None: continue
    D_Sisaku[str(a)] = (ws.cell(r,3).value, ws.cell(r,4).value)

# D_kumitori: 組取り数 A→B
D_kumitori = {}
ws = wbv['組取り数']
for r in range(2, ws.max_row+1):
    a = ws.cell(r,1).value
    if a is None: continue
    D_kumitori[str(a)] = ws.cell(r,2).value

# ChkGroup: 試作 A→B を先に、次に グループ B(品目)→A(グループ)
ChkGroup = {}
ws = wbv['試作']
for r in range(2, ws.max_row+1):
    a, b = ws.cell(r,1).value, ws.cell(r,2).value
    if a and str(a) not in ChkGroup: ChkGroup[str(a)] = str(b) if b else ""
ws = wbv['グループ']
for r in range(2, ws.max_row+1):
    g, it = ws.cell(r,1).value, ws.cell(r,2).value
    if it and str(it) not in ChkGroup: ChkGroup[str(it)] = str(g) if g else ""

# BOSS読み込み
ws = wbv['BOSS']
boss = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] is None and row[1] is None: continue
    boss.append(row)

# カスタム設定 日程変更(A品目,B変更前,C変更後 / R指図,S変更後) をBOSS N列へ適用
D_tmp = {}
ws = wbv['カスタム設定']
for r in range(2, ws.max_row+1):
    a, b, c = ws.cell(r,1).value, ws.cell(r,2).value, ws.cell(r,3).value
    if a and hasattr(b,'year') and hasattr(c,'year'):
        key = f"{a}|{b.strftime('%Y/%m/%d')}"
        if key not in D_tmp: D_tmp[key] = c
R_tmp = {}
for r in range(2, ws.max_row+1):
    s, t = ws.cell(r,18).value, ws.cell(r,19).value
    if s and hasattr(t,'year'):
        R_tmp[str(s)] = t

new_boss = []
for row in boss:
    row = list(row)
    it = row[1]; nd = row[13]
    key = f"{it}|{nd.strftime('%Y/%m/%d')}" if hasattr(nd,'year') and it else None
    if key and key in D_tmp:
        row[13] = D_tmp[key]
    if str(row[0]) in R_tmp:
        row[13] = R_tmp[str(row[0])]
    new_boss.append(row)
boss = new_boss

# 指図タイプフィルタ ^(LA|ZP1\d|ZP2\d)$ かつ KV開始品目除外
reg = re.compile(r'^(LA|ZP1\d|ZP2\d)$')
boss = [r for r in boss if r[3] and reg.match(str(r[3])) and not str(r[1]).startswith('KV')]

# D_Day: 日付→{品目:数量}
D_Day = defaultdict(lambda: defaultdict(float))
for row in boss:
    nd = row[13]
    if hasattr(nd,'year'):
        d = nd.date()
        q = row[11] if isinstance(row[11],(int,float)) else 0
        D_Day[d][str(row[1])] += q

def chkTime(D_items):
    """D_items: {品目: 数量} → (量産h, NotFound, 試作h)"""
    x = 0.0; x_sisaku = 0.0; kotei = 0.0
    NotFound = []
    for it, qty in D_items.items():
        sisaku = it in D_Sisaku
        if sisaku:
            c = D_Sisaku[it][0]
            if isinstance(c,(int,float)): kotei += c
        ct = D_CycleTime.get(it)
        ct_ok = isinstance(ct,(int,float))
        if ct_ok:
            if sisaku: x_sisaku += qty*ct + SETUP
            else:      x += qty*ct + SETUP
        else:
            if it in D_kumitori:
                k = D_kumitori[it]
                if sisaku: x_sisaku += qty/k*KARI_SISAKU + SETUP
                else:      x += qty/k*KARI + SETUP
                tag = it
            else:
                if sisaku:
                    ks = D_Sisaku[it][1]
                    if isinstance(ks,(int,float)) and ks != "":
                        x_sisaku += qty/ks*KARI_SISAKU + SETUP
                        tag = it
                    else:
                        x_sisaku += qty*KARI_SISAKU + SETUP
                        tag = "★"+it
                else:
                    x += qty*28 + SETUP
                    tag = "★"+it
            NotFound.append(tag)
    return x/3600/EFF, "|".join(NotFound), kotei + x_sisaku/3600/EFF

# 計算シートの現状値を読み取り比較
wsc = wbv['計算']
cal_dates = []
for r in range(2, wsc.max_row+1):
    d = wsc.cell(r,1).value
    if hasattr(d,'year'):
        cal_dates.append((r, d.date()))

headers = {}
for c in range(1, 22):
    h = wsc.cell(1,c).value
    if h: headers[str(h)] = c

diffs = [] ; matches = 0 ; total = 0
recomputed = {}   # (date) -> {group: hours}
notfound_map = {}
for r, d in cal_dates:
    items = D_Day.get(d, {})
    groups = defaultdict(dict)
    for it, q in items.items():
        g = ChkGroup.get(it, "フリー")
        groups[g][it] = q
    rec = {}
    for g, dmap in groups.items():
        mh, nf, sh = chkTime(dmap)
        val = sh if sh > 0 else mh
        rec[g] = val
        if nf: notfound_map.setdefault(d, []).extend(nf.split("|"))
    recomputed[d] = rec
    # 比較(値が存在する列のみ)
    for g, hv in rec.items():
        c = headers.get(g)
        if c is None: continue
        stored = wsc.cell(r, c).value
        total += 1
        if isinstance(stored,(int,float)) and abs(stored-hv) < 1e-6:
            matches += 1
        else:
            diffs.append((d, g, hv, stored))

print(f"比較セル数={total} 一致={matches} 不一致={len(diffs)}")
for d,g,hv,st in diffs[:25]:
    print(f"  不一致 {d} {g}: 再現={hv:.4f} 保存値={st}")

# 保存値側にあって再現にない列(逆方向)
rev = []
for r, d in cal_dates:
    for g, c in headers.items():
        if g in ('日付','DBにBOM無し','DBにサイクルタイム無し','フリー容量','フリー部品種類数'): continue
        stored = wsc.cell(r,c).value
        if isinstance(stored,(int,float)):
            if g not in recomputed.get(d,{}):
                rev.append((d,g,stored))
print(f"\n保存値のみ(再現ゼロ)セル数={len(rev)}")
for d,g,v in rev[:25]:
    print(f"  {d} {g}={v:.4f}")

wbv.close(); wbf.close()
