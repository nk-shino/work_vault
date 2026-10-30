# -*- coding: utf-8 -*-
"""サーバー管理手順書.docx の構造抽出(読み取り専用)"""
import sys, io, zipfile, re
from collections import Counter
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = r"C:\Users\ss80751\OneDrive - NK\デスクトップ\サーバー管理手順書.docx"
OUT = r"C:\Users\ss80751\kohden-code-workspace\20261030_サーバー管理手順書解析\_structure.txt"

buf = io.StringIO()
w = buf.write

d = docx.Document(SRC)

# --- core properties ---
cp = d.core_properties
w("=== core properties ===\n")
for k in ("title", "author", "last_modified_by", "created", "modified", "subject", "comments"):
    v = getattr(cp, k, None)
    w(f"{k}: {v}\n")

# --- sections ---
w("\n=== sections ===\n")
for i, s in enumerate(d.sections):
    w(f"[{i}] page {s.page_width.cm:.1f}x{s.page_height.cm:.1f}cm "
      f"margin T{s.top_margin.cm:.2f} B{s.bottom_margin.cm:.2f} "
      f"L{s.left_margin.cm:.2f} R{s.right_margin.cm:.2f} "
      f"orient={'LANDSCAPE' if s.orientation is not None and str(s.orientation).endswith('LANDSCAPE') else 'portrait'}\n")
    for label, part in (("header", s.header), ("footer", s.footer)):
        txts = [p.text for p in part.paragraphs if p.text.strip()]
        if txts:
            w(f"  {label}: {' / '.join(t.strip() for t in txts)[:120]}\n")

# --- styles histogram ---
w("\n=== paragraph style histogram ===\n")
stc = Counter()
for p in d.paragraphs:
    stc[p.style.name if p.style is not None else "?"] += 1
for name, n in stc.most_common():
    w(f"{n:5d}  {name}\n")

# --- default fonts (styles.xml) ---
w("\n=== fonts in styles.xml (rFonts) ===\n")
with zipfile.ZipFile(SRC) as z:
    styles = z.read("word/styles.xml").decode("utf-8", "ignore")
fonts = Counter(re.findall(r'w:(?:ascii|eastAsia)="([^"]+)"', styles))
for f, n in fonts.most_common(12):
    w(f"{n:5d}  {f}\n")
    break
w("default:\n")
m = re.search(r'<w:rPrDefault>.*?</w:rPrDefault>', styles, re.S)
if m:
    for f in re.findall(r'w:(?:ascii|eastAsia|hAnsi)="([^"]+)"', m.group(0)):
        w(f"  {f}\n")

# --- body outline in order ---
w("\n=== body outline (headings only) ===\n")
body = d.element.body
idx_img = 0
def para_images(p_el):
    return len(p_el.findall(".//" + qn("w:drawing"))) + len(p_el.findall(".//" + qn("w:pict")))

n_para = n_tbl = n_img = 0
for child in body.iterchildren():
    if child.tag == qn("w:p"):
        p = Paragraph(child, d)
        n_para += 1
        n_img += para_images(child)
        st = p.style.name if p.style is not None else ""
        if re.match(r"(Heading|見出し)", st) and p.text.strip():
            w(f"{st}: {p.text.strip()[:80]}\n")
    elif child.tag == qn("w:tbl"):
        n_tbl += 1
w(f"\nparas={n_para} tables={n_tbl} images(drawing/pict)={n_img} inline_shapes={len(d.inline_shapes)}\n")

# --- full body dump with style + table contents ---
w("\n=== full body dump ===\n")
ti = 0
for child in body.iterchildren():
    if child.tag == qn("w:p"):
        p = Paragraph(child, d)
        st = p.style.name if p.style is not None else "?"
        t = p.text.strip()
        im = para_images(child)
        num = "N" if child.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None and child.find(".//" + qn("w:numPr")) is not None else ""
        flag = ""
        if im:
            flag = f" [IMGx{im}]"
        if t or flag:
            w(f"<{st}>{num} {t[:200]}{flag}\n")
    elif child.tag == qn("w:tbl"):
        ti += 1
        tb = Table(child, d)
        rows = len(tb.rows)
        cols = len(tb.columns)
        w(f"<TABLE {ti}> {rows}x{cols}\n")
        for ri, row in enumerate(tb.rows[:3]):
            cells = [c.text.strip().replace("\n", "⏎")[:40] for c in row.cells]
            w(f"  r{ri}: {' | '.join(cells)[:300]}\n")
        if rows > 3:
            w(f"  ... {rows-3} more rows\n")

with open(OUT, "w", encoding="utf-8") as f:
    f.write(buf.getvalue())
print(f"written {OUT}, {len(buf.getvalue())} chars")
