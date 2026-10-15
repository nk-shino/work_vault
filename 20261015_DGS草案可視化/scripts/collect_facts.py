# -*- coding: utf-8 -*-
"""根拠収集の最終スクリプト: 4対象Excelの実測一覧を JSON として出力する(原本は読み取り専用)"""
import openpyxl, xlrd, zipfile, re, json, os

BASE = r"C:\Users\ss80751\OneDrive - NK\ドキュメント\Work_Vault\10_業務全般\PanaCIM_LNB"
OUT = r"C:\Users\ss80751\kohden-code-workspace\20261015_DGS草案可視化\evidence"
os.makedirs(OUT, exist_ok=True)

result = {}

# --- 1. DGS P1.xls ---
p1 = os.path.join(BASE, r"02_参照元\マウントプログラム(DGS)\0494-913473 実装プログラム手順(DGS)_P1.xls")
wb = xlrd.open_workbook(p1)
sh = wb.sheet_by_name("付表1")
p1_cells = {}
for r in range(sh.nrows):
    for c in range(sh.ncols):
        v = sh.cell_value(r, c)
        if isinstance(v, str) and v.strip() and v.strip() not in ("  ×",):
            p1_cells[f"r{r+1}c{c+1}"] = v.strip()[:60]
result["P1"] = {"sheets": wb.sheet_names(), "cells": p1_cells}

# --- 2. DGS P2.xlsx ---
p2 = os.path.join(BASE, r"02_参照元\マウントプログラム(DGS)\0494-913473 実装プログラム手順(DGS)_P2.xlsx")
wb2 = openpyxl.load_workbook(p2, read_only=False)
sheets = []
for ws in wb2.worksheets:
    o = ws.page_setup.orientation
    pa = bool(ws.print_area)
    rb = len(ws.row_breaks.brk) if ws.row_breaks else 0
    foot = (ws.oddFooter.center.text or "") if ws.oddFooter else ""
    sheets.append({"name": ws.title, "rows": ws.max_row, "cols": ws.max_column,
                   "orientation": o, "print_area": pa, "row_breaks": rb, "footer": foot})
result["P2"] = {"sheets": sheets, "count": len(sheets)}
# フロー図の遷移先(Z列)
z2 = {}
for s in ["試作フロー", "設計変更_部品フロー", "設計変更_改版フロー", "新製品フロー"]:
    ws = wb2[s]
    items = []
    for row in ws.iter_rows(min_col=26, max_col=26):
        for c in row:
            if c.value and "➡" in str(c.value):
                items.append({"cell": c.coordinate, "text": str(c.value)[:40]})
    z2[s] = items
result["P2"]["transitions"] = z2
# パス掲載シート
path_sheets = []
for ws in wb2.worksheets:
    hit = False
    for row in ws.iter_rows(max_col=6):
        for c in row:
            if c.value and "nksv" in str(c.value).lower():
                hit = True
                break
        if hit:
            break
    if hit:
        path_sheets.append(ws.title)
result["P2"]["path_sheets"] = path_sheets
# 数式($G$4参照)
fx = []
for name in ["パーツ_マウンタ設定", "パーツ_その他設定"]:
    ws = wb2[name]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                fx.append({"sheet": name, "cell": c.coordinate, "formula": c.value[:20]})
result["P2"]["formulas"] = fx
# 印刷範囲設定シート
result["P2"]["print_area_sheets"] = [ws.title for ws in wb2.worksheets if ws.print_area]
# DGS起動の認証情報セル
result["P2"]["secret_cell"] = {"sheet": "DGS起動", "cell": "A12", "masked": True}
# hyperlink
z = zipfile.ZipFile(p2)
hl = 0
for n in z.namelist():
    if "worksheets/_rels" in n:
        hl += z.read(n).decode("utf-8", "replace").count("hyperlink")
result["P2"]["hyperlinks"] = hl
# 図形
dr = [n for n in z.namelist() if re.match(r"xl/drawings/drawing\d+\.xml$", n)]
sp = cx = pic = 0
for n in dr:
    d = z.read(n).decode("utf-8", "replace")
    sp += d.count(":sp>")
    cx += d.count(":cxnSp>")
    pic += d.count(":pic>")
result["P2"]["shapes"] = {"drawing_files": len(dr), "sp": sp, "cxn": cx, "pic": pic}
# DGS起動の画像アンカー行
d = z.read("xl/drawings/drawing1.xml").decode("utf-8", "replace") if "xl/drawings/drawing1.xml" in z.namelist() else ""
rows = re.findall(r"<xdr:row>(\d+)</xdr:row>", d)
result["P2"]["dgs_start_image_rows"] = rows

# --- 3. PanaCIM 会議版 ---
pc = os.path.join(BASE, "PanaCIM管理手順書_草案(会議中に手を加えたもの).xlsx")
wb3 = openpyxl.load_workbook(pc)
sheets3 = []
for ws in wb3.worksheets:
    b1 = str(ws["B1"].value or "")
    sheets3.append({"name": ws.title, "rows": ws.max_row, "cols": ws.max_column,
                    "B1_memo": (b1[:40] if b1 not in ("None", "") else None),
                    "landscape": ws.page_setup.orientation == "landscape",
                    "print_area": bool(ws.print_area)})
# 目次リンクと切れ
toc_links = []
sheetset = set(wb3.sheetnames)
ws = wb3["目次"]
for row in ws.iter_rows():
    for c in row:
        if c.hyperlink:
            t = c.hyperlink.target or ""
            m = re.match(r"#'([^']+)'!", t)
            nm = m.group(1) if m else t
            toc_links.append({"cell": c.coordinate, "label": str(c.value), "target": nm,
                              "ok": nm in sheetset})
result["PanaCIM"] = {"sheets": sheets3, "toc_links": toc_links,
                     "count": len(wb3.sheetnames)}
# 表紙
ws = wb3["表紙"]
result["PanaCIM"]["cover"] = {f"r{r}": str(ws.cell(r, 2).value)[:44] for r in range(5, 16) if ws.cell(r, 2).value}
# 改訂履歴の最終版
ws = wb3["改訂履歴"]
last = None
for r in range(9, ws.max_row + 1):
    if ws.cell(r, 1).value:
        last = [str(ws.cell(r, c).value)[:40] for c in range(1, 4)]
result["PanaCIM"]["last_rev"] = last
# 申し送り相当のメモ一覧(B1メモ6件)
memos = [(s["name"], s["B1_memo"]) for s in sheets3 if s["B1_memo"]]
result["PanaCIM"]["B1_memos"] = memos

# --- 4. LNB 会議版 ---
lb = os.path.join(BASE, "LNB管理手順書_草案(会議中に手を加えたもの).xlsx")
wb4 = openpyxl.load_workbook(lb)
sheets4 = []
for ws in wb4.worksheets:
    b1 = str(ws["B1"].value or "")
    sheets4.append({"name": ws.title, "rows": ws.max_row, "cols": ws.max_column,
                    "B1_memo": (b1[:40] if b1 not in ("None", "") else None),
                    "landscape": ws.page_setup.orientation == "landscape",
                    "print_area": bool(ws.print_area)})
sheetset4 = set(wb4.sheetnames)
ws = wb4["目次"]
toc_links4 = []
for row in ws.iter_rows():
    for c in row:
        if c.hyperlink:
            t = c.hyperlink.target or ""
            m = re.match(r"#'([^']+)'!", t)
            nm = m.group(1) if m else t
            toc_links4.append({"cell": c.coordinate, "label": str(c.value), "target": nm,
                               "ok": nm in sheetset4})
ws = wb4["表紙"]
cover4 = {f"r{r}": str(ws.cell(r, 2).value)[:44] for r in range(5, 16) if ws.cell(r, 2).value}
ws = wb4["改訂履歴"]
last4 = None
for r in range(9, ws.max_row + 1):
    if ws.cell(r, 1).value:
        last4 = [str(ws.cell(r, c).value)[:40] for c in range(1, 4)]
result["LNB"] = {"sheets": sheets4, "toc_links": toc_links4, "cover": cover4,
                 "last_rev": last4, "count": len(wb4.sheetnames)}

with open(os.path.join(OUT, "excel_facts.json"), "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=1)
print("saved", os.path.join(OUT, "excel_facts.json"))
print("P1 cells:", len(p1_cells), "| P2 sheets:", len(sheets),
      "| PanaCIM:", len(sheets3), "toc_links:", len(toc_links),
      "| LNB:", len(sheets4), "toc_links:", len(toc_links4))
