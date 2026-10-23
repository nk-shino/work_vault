# -*- coding: utf-8 -*-
"""固定テーブル棚卸し 作業手順書（統合版）.docx を作成する。

出典:
  - スライド1-20.JPG（佐藤勇希氏「固定テーブル 棚卸し方法」）
  - 固定テーブル棚卸し方法_手順書_v2.md（2026-09-15）
  - 棚卸しメモ「システム周り手順」
  - 0確認マクロver2_運用手順.md（2026-10-22）
新規の判断は加えない。既存資料の統合。
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SLIDE_DIR = r"C:\Users\ss80751\OneDrive - NK\ドキュメント\Work_Vault\07-1_棚卸\原典資料\固定テーブル棚卸し方法_USB版"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "固定テーブル棚卸し作業手順書_2026-10-23.docx")

JP_FONT = "游ゴシック"      # eastAsia
LAT_FONT = "Yu Gothic"      # ascii/hAnsi

doc = Document()

# ---- page setup: A4 ----
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)

# ---- base style ----
style = doc.styles["Normal"]
style.font.name = LAT_FONT
style.font.size = Pt(10.5)
rpr = style.element.get_or_add_rPr()
rf = rpr.find(qn("w:rFonts"))
if rf is None:
    rf = OxmlElement("w:rFonts")
    rpr.append(rf)
rf.set(qn("w:ascii"), LAT_FONT)
rf.set(qn("w:hAnsi"), LAT_FONT)
rf.set(qn("w:eastAsia"), JP_FONT)
pf = style.paragraph_format
pf.space_after = Pt(4)
pf.line_spacing = 1.15


def set_eastasia(run):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), LAT_FONT)
    rFonts.set(qn("w:hAnsi"), LAT_FONT)
    rFonts.set(qn("w:eastAsia"), JP_FONT)


def para(text="", size=10.5, bold=False, color=None, align=None,
         space_before=0, space_after=4, indent=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    set_eastasia(r)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def h1(text):
    para(text, size=14, bold=True, color="1F4E78", space_before=14, space_after=6)


def h2(text):
    para(text, size=12, bold=True, color="2E5E8C", space_before=10, space_after=4)


def step(text):
    """手順の1行。番号付きテキスト行。"""
    para(text, indent=0.4, space_after=2)


def note(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string("555555")
    set_eastasia(r)
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(6)


def slide(n, width_cm=15.5):
    path = os.path.join(SLIDE_DIR, f"スライド{n}.JPG")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    cap = para(f"スライド{n}", size=9, align="center", color="777777", space_after=8)
    return cap


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def fill(cell, text, bold=False, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    set_eastasia(r)
    p.paragraph_format.space_after = Pt(1)


def table(rows, header=True, widths=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            fill(c, val, bold=(header and i == 0))
            if header and i == 0:
                shade(c, "DCE6F1")
            if widths:
                c.width = Cm(widths[j])
    if widths:
        for j, w in enumerate(widths):
            for i in range(len(rows)):
                t.cell(i, j).width = Cm(w)
    return t


# ================= 表紙部 =================
para("固定テーブル棚卸し 作業手順書（統合版）", size=20, bold=True, align="center", space_before=24, space_after=8)
para("2026-10-23 作成（第1版）", size=11, align="center", space_after=16)

para("対象製品", bold=True, space_before=6)
para("AED・TEC1000・TEC5600・NKV・送信機（5製品）", indent=0.4)

para("この文書の出典", bold=True, space_before=6)
step("①「固定テーブル 棚卸し方法」（佐藤勇希氏・スライド20枚）— 本文中の画像はこの資料から転載")
step("②固定テーブル棚卸し方法_手順書_v2（2026-09-15）")
step("③棚卸しメモ「システム周り手順」（EG-Pro・SQLまわりの操作）")
step("④0確認マクロver2_運用手順（2026-10-22）・棚卸マクロver2 の進行状況（2026-10-15）")

para("位置づけ", bold=True, space_before=6)
para("既存資料を統合したもので、新たな判断は加えていない。手順の途中で分からない点が出た場合は、"
     "その場で判断せず現場責任者に確認する。", indent=0.4)

doc.add_page_break()

# ================= 0. 使い方 =================
h1("0. この手順書の使い方")
para("手順の中で「〈製品名〉」と書いた箇所は、その日操作する設備の製品名に読み替える。倉庫やSQLで使うA55コードは下の表のとおり。")
table([
    ["製品名", "A55コード", "LINE"],
    ["NKV", "201", "LINE1"],
    ["TEC5600 / TEC1000 / 送信機", "202", "LINE2"],
    ["AED", "203", "LINE3"],
], widths=[7.5, 4.0, 3.0])
note("倉庫名の正式な表記は EG-Pro の倉庫検索画面で確認する。A55コードのうち213は、事前準備で使えることを確認する範囲。")

# ================= 1. 目的 =================
h1("1. 目的")
para("固定テーブル棚卸しは、帳票上のリールの所在とステータスを、現場の実物と一致させる作業。"
     "在庫数そのものを直接修正する作業は含まない。数値に矛盾が見つかった場合は、"
     "棚卸し完了後に別途確認して修正する。")

# ================= 2. 取り消せない操作 =================
h1("2. 取り消せない操作（作業前に読む）")
step("① idouSQL（倉庫移動SQL）は COMMIT の後に自動では戻せない。実行前に SELECT 結果を画面で目視確認する。"
     "表示上のR列・T列の値は信用しない。")
step("② 0確認マクロによるラベル消去は取り消せない。実行前に対象リストを確認する。")
step("③ リール破棄処理も同様に取り消せない。対象の確認を必ず行う。")
step("④ 当日のダミーJob（例: 000001-003-DMY など ****-DMY の形式）が作成済みか、作業の最初に確認する。")

# ================= 3. 全体の流れ =================
h1("3. 全体の流れ（早見表）")
table([
    ["No", "工程", "主な内容", "スライド"],
    ["Step1", "設備操作", "Job読込・切替（1人で作業できる）", "2"],
    ["Step2", "セットアップレポート出力", "Job読込10分後にeplantviewで取得・保存", "3-4"],
    ["Step3", "EG-Pro在庫出力", "方法1: cgi URL / 方法2: SSMSクエリ", "5-8"],
    ["Step4", "マクロ実行（配列データ投入）", "配列情報入力・対応表・固定グループ名・失敗リスト", "9"],
    ["Step5", "スプライシング切替失敗処理", "失敗リールの解除と再出力", "10"],
    ["Step6", "カウント前準備", "ラベル貼り・仮止め・新リールカウンタ", "11-12"],
    ["Step7", "カウント実施", "バーコード読取・要確認BOX", "13-14"],
    ["Step8", "入庫処理", "ダミーJob→手動入庫または倉庫移動SQL", "15-17"],
    ["Step9", "所在不明リール処理", "所在不明リスト・分岐判断", "18"],
    ["Step10", "部品取り付け", "取付・例外分岐の確認", "19-20"],
], widths=[1.6, 5.0, 7.6, 2.2])

doc.add_page_break()

# ================= Step1 =================
h1("Step1 設備操作")
step("① 設備側のボタン操作（図1・図2）でJob読込と切替を行う。1人で作業できる。")
step("② 読み込んだJobが設備に反映されていることを確認してから次の工程へ。")
slide(2)

# ================= Step2 =================
h1("Step2 セットアップレポート出力")
step("① Job読込の後、約10分待つ。EG-Proへの反映待ちの時間。")
step("② eplantview で「材料照合」→「現在のセットアップレポート」を選択する。")
step("③ 設備選択で「全設備」→OK→適用→保存の順に進める。")
step("④ 保存したファイルは Shift-JIS で出る。このファイルを Step4 のマクロに読み込ませる。")
slide(3)
slide(4)

doc.add_page_break()

# ================= Step3 =================
h1("Step3 EG-Pro在庫情報出力")
h2("方法1（cgi・推奨）")
step("① EG-Pro の cgi URL にアクセスして在庫情報を取得する。事務PCからでも実行できる。")
h2("方法2（SSMS）")
step("① SQL Server Management Studio で EG-Pro のデータベースに接続する。")
step("② 「在庫データ取得.txt」のクエリを貼り付けて実行する。")
step("③ 実行結果を保存する。")
para("どちらか一方でよい。取得した在庫データは Step4 のマクロに読み込ませる。", space_before=4)
slide(5)
slide(6)

doc.add_page_break()

h2("方法2の続き")
slide(7)
slide(8)

doc.add_page_break()

# ================= Step4 =================
h1("Step4 マクロ実行（配列データ投入）")
step("① 配列情報入力マクロを開き、Step2 のセットアップレポートと Step3 の在庫データを指定して実行する。")
step("② A55コードの対応表と固定グループ名はマスタに固定されている。変更が必要な場合は、"
     "現場責任者に確認してから直す。")
step("③ 処理が終わると失敗リストが出る。失敗行は Step5 の手順へ進む。")
slide(9)

# ================= Step5 =================
h1("Step5 スプライシング切替失敗処理")
step("① 失敗リストに載ったリールに対し、切替解除の手順（3工程）を行う。")
step("② 解除できたら、Step3 の出力から Step4 のマクロ実行までをやり直す。失敗がなくなるまで繰り返す。")
step("③ スプライシング部品リストを印刷しておく。Step7・Step10 で使う。")
slide(10)

doc.add_page_break()

# ================= Step6 =================
h1("Step6 カウント前準備（ラベル貼り）")
step("① テーブルからリールを外し、赤枠で示したシリアルラベルを再発行して仮止めする。")
step("② 青丸シールを貼る。新リールカウンタを用意する。")
slide(11)
slide(12)

doc.add_page_break()

# ================= Step7 =================
h1("Step7 カウント実施")
step("① バーコードを読む。リールにバーコードが1枚のときと2枚のときで、読み方が変わる（画像参照）。")
step("② 巻き戻しはしない。読めたら次のリールへ進む。")
step("③ 「要確認BOX」: バーコード1枚なのにスプライシング確認画面が出る場合、"
     "2枚なのに確認画面が出ない場合は確認対象。扱いは Step10 の分岐と同じなので、"
     "その場で判断せず記録して進む。")
slide(13)
slide(14)

doc.add_page_break()

# ================= Step8 =================
h1("Step8 入庫処理")
h2("① ダミーJobの操作")
step("EG-Pro で ****-DMY 形式のダミーJobを選択→「生産開始」→すぐに「キャンセル」。"
     "これでリールがハズシ(Lx)に切り替わる。")
slide(15)
h2("② 手動入庫（バーコード作成→ハンディ読取）")
step("手動入庫用のバーコードを作成し、ハンディで読む。読む順序は上→下。")
note("2025-03-21追記（スライド16記載のとおり）: 倉庫入庫ではなく、棚卸し用倉庫（A55 201/202/203）を指定する。"
     "ロケーション移動でも対応できる。")
slide(16)
h2("③ ロケーション移動（SQL）")
step("移動先を入力→現場部品リスト作成→倉庫移動SQL作成→実行。実行しても画面上は何も起きていないように見える。"
     "EG-Pro の倉庫検索で反映を確認する。")
step("実行前の SELECT 結果の目視確認が唯一の防御になる（「2. 取り消せない操作」参照）。")
slide(17)

doc.add_page_break()

# ================= Step9 =================
h1("Step9 所在不明リール処理")
step("① 全ての入庫が終わったら、所在不明リストを出す。")
step("② リール消費推奨マクロで候補を出し、「0確認」「ライン倉庫」「ラベル発行」「借用」のいずれかに分岐する。")
step("③ 判断に迷うものは現場責任者へ回す。")
slide(18)

# ================= Step10 =================
h1("Step10 部品取り付け")
step("① バーコードが1個のときと2個のときで手順が分かれる（画像参照）。")
step("② テープ長を確認し、スプライシング済みの確認メッセージが出ることを確認してから取り付ける。")
step("③ テープ長が不足しているときは、切替を解除してリール本体のバーコードを読む。")
step("④ 例外分岐: 1枚なのに確認画面が出た場合はテープの有無で判断する。"
     "2枚なのに確認画面が出ない場合は、棚卸し完了の前後で扱いが変わる（画像参照）。")
slide(19)
slide(20)

doc.add_page_break()

# ================= ツールの現状 =================
h1("補足1. ツールの現状（ver2、2026-10-23時点）")
h2("0確認マクロ")
step("・ 正式版は 0確認マクロ_ver2.xlsm。保存場所は共有 \\\\nksv569.ad.nkc.co.jp\\Share\\mnt\\DGS-SV\\◎棚卸し。")
step("・ 運用の規律は 0確認マクロver2_運用手順.md に確定している（現場責任者の回答 Q1〜Q35 に基づく。"
     "Q14・Q35-C の実装と試験は2026-10-22に合格）。")
step("・ 実行者は現場責任者と、その指名を受けた担当者。初回実行と、正常運用の確認が取れるまでの実行には立会いが必要。")
h2("棚卸マクロ ver2（modTanaoroshi.bas）")
step("・ コンパイルの通過は確認済み（2026-10-15）。実機での試し実行はまだ行っていない。")
step("・ 当日の本番は従来どおり現行マクロで運用し、ver2 は試し実行モードで並行して検証する。"
     "ver2 の使用は、現場責任者の指示がある場合に限る。")

# ================= 失敗とリカバリー =================
h1("補足2. 失敗とリカバリー（抜粋）")
table([
    ["状況", "対応"],
    ["セットアップレポートにJobが反映されていない", "10分待ってから再取得する（Step2 ①）"],
    ["配列マクロで失敗リストが出る", "Step5 の解除手順を行ってから再実行する"],
    ["スプライシング切替失敗のリールが残る", "解除→再出力を繰り返す。残った分は報告する"],
    ["バーコードが読めない（テープ損傷など）", "リール本体のバーコードか手動読取。判断は現場責任者に確認"],
    ["idouSQL 実行後に数が合わない", "COMMIT 前なら ROLLBACK。COMMIT 後は報告する（自動では戻せない）"],
    ["所在不明のリールが出る", "Step9 の分岐。迷うものは現場責任者へ"],
], widths=[7.5, 9.0])
note("一覧は抜粋。詳細は 固定テーブル棚卸し方法_手順書_v2 の失敗とリカバリーの節を参照。")

# ================= チェックリスト =================
h1("補足3. 自己チェックリスト（終了前の確認）")
for item in [
    "□ 当日のダミーJobは作成・確認したか",
    "□ セットアップレポートは Job読込 10分後のものか",
    "□ 在庫データは取得できたか（方法1/2のどちらか）",
    "□ 失敗リストの処理は終わったか（部品リストの印刷物はあるか）",
    "□ 要確認BOXに残ったものを記録したか",
    "□ 入庫の後、EG-Pro の倉庫検索で所在の反映を確認したか",
    "□ idouSQL 実行前に SELECT 結果を目視確認したか",
    "□ 所在不明リストを出し、分岐の判断をしたか",
    "□ 部品取り付けでスプライシング済みの確認メッセージを見たか",
    "□ 例外作業の記録（日時・対象・判断者）を残したか",
]:
    step(item)

# ================= 報告 =================
h1("補足4. 現場から事務・現場責任者へ報告するもの")
step("・ 要確認BOXに残ったもの")
step("・ バーコード2枚なのに確認画面が出ないもの")
step("・ 失敗リストで解決しなかったもの")
step("・ 所在不明のまま残ったリール")
step("・ 例外作業を行った記録")

para("— 以上 —", align="center", space_before=16)
note("本書の画像は佐藤勇希氏の資料（固定テーブル 棚卸し方法）から転載。手順の最終確認は、現場の設備と "
     "EG-Pro の実画面で行うこと。")

doc.save(OUT)
print("saved:", OUT)
