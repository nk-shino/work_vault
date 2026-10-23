# -*- coding: utf-8 -*-
"""7-4 検証v3: 21手順+発動条件の原票突合(パスワードはD149から動的抽出)。"""
import sys, io, re, warnings
warnings.filterwarnings("ignore")
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

SRC = r"C:\Users\ss80751\OneDrive - NK\ドキュメント\Work_Vault\10_業務全般\PanaCIM_LNB\02_参照元\様式_操作手順書0494-912039\0494-912039(06)-P2-P47.xlsx"
DST = r"C:\Users\ss80751\OneDrive - NK\ドキュメント\Work_Vault\10_業務全般\PanaCIM_LNB\01_成果物\20261015_管理手順書案\Pana-CIM管理手順書案_未承認.xlsx"
RED = "［認証情報：伏字］"

STEP_MAP = [
    ("C42", "B9"), ("C43", "B11"), ("C44", "D12"), ("C45", "B31"), ("C46", "B60"),
    ("C47", "D61"), ("C48", "B82"), ("C49", "D83"), ("C50", "B114"), ("C51", "B128"),
    ("C52", "D129"), ("C53", "B148"), ("C54", "D149"), ("C55", "D151"), ("C56", "B169"),
    ("C57", "B180"), ("C58", "B201"), ("C59", "B223"), ("C60", "B235"), ("C61", "D253"),
    ("C62", "B277"),
]

ssrc = openpyxl.load_workbook(SRC, read_only=True)["PanaCIM"]
def sv(a):
    v = ssrc[a].value
    return v if isinstance(v, str) else ("" if v is None else str(v))

d149 = sv("D149")
m = re.search(r"「([^」]+)」", d149)
pw = m.group(1) if m else None
assert pw, "パスワード抽出失敗"

dst = openpyxl.load_workbook(DST)["異常時の対応"]

ok = 0
for dst_addr, src_addr in STEP_MAP:
    exp = sv(src_addr).replace(pw, RED)
    d = dst[dst_addr].value or ""
    match = (d == exp)
    ok += match
    if not match:
        print(f"{dst_addr}←{src_addr}: NG")
        print(f"  期待: {exp!r}")
        print(f"  実際: {d!r}")
print(f"手順突合: {ok}/{len(STEP_MAP)}")

c28 = dst["C28"].value or ""
print("発動条件含有: B6=", sv("B6") in c28, " B7=", sv("B7") in c28, " C7=", sv("C7") in c28)

# パスワード原文字列が領域に残っていないか
leak = [c.coordinate for row in dst.iter_rows(min_row=1, max_row=200, min_col=1, max_col=60)
        for c in row if isinstance(c.value, str) and pw in c.value]
print("パスワード残存:", leak if leak else "なし")

# QC行の主張と突合結果の整合チェック
c36 = dst["C36"].value or ""
n_red = sum(1 for _, a in STEP_MAP if RED in (dst[_].value or ""))
print(f"伏字セル数={n_red} (C36の主張: {c36})")
